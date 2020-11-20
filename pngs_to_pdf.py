import os
import cv2
import numpy as np
from docx2pdf import convert
from docx import Document
from docx.shared import Inches, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# 先清空圖片暫存資料夾
for i in os.listdir(r'C:\Users\user\Desktop\demo_elastic\cv'):
    os.remove(r'C:\Users\user\Desktop\demo_elastic\cv\%s' % i)

# 取出之圖表，每個圖新增一組資料，分別為[來源圖編號, X1, X2, Y1, Y2]
vis_list = [[1, 6, 452, 63, 484], [4, 2, 639, 55, 585], [6, 4, 1267, 61, 585]]

# 從截圖取出圖表
for i in range(len(vis_list)):
    src = cv2.imdecode(np.fromfile(r"C:\Users\user\Desktop\demo_elastic\screenshot%s.png" % vis_list[i][0], dtype=np.uint8), -1)
    src = src[vis_list[i][3]:vis_list[i][4], vis_list[i][1]:vis_list[i][2]] # Y1:Y2, X1:X2
    cv2.imencode('.jpeg', src)[1].tofile(r"C:\Users\user\Desktop\demo_elastic\cv\cv%s.jpg" % i)

# 將圖表貼到Word
doc = Document()  # doc对象
section = doc.sections[0]
section.left_margin = Cm(1.27)
section.right_margin = Cm(1.27)
section.top_margin = Cm(1.27)
section.bottom_margin = Cm(1.27)

for i in range(len(vis_list)):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = paragraph.add_run()
    run.add_picture(r"C:\Users\user\Desktop\demo_elastic\cv\cv%s.jpg" % i, width=Inches(6.7))  # 添加图, 设置宽度

doc.save(r'C:\Users\user\Desktop\demo_elastic\test.docx')  # 保存路径
convert(r'C:\Users\user\Desktop\demo_elastic\test.docx',r'C:\Users\user\Desktop\demo_elastic\test.pdf')
os.remove(r'C:\Users\user\Desktop\demo_elastic\test.docx')
